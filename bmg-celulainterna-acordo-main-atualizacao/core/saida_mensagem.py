"""
Writer da planilha de saida da Tarefa 1 (mensagens de acordo).

Schema:
  CIV | Cliente | CNJ | Estado | Mensagem | Data_Pagamento | Saldo_Atualizado | Capturado_em

Data_Pagamento e Saldo_Atualizado vem da aba Financeiro (lancamento
mais recente). Sao opcionais: ficam vazios se a aba Financeiro nao
tiver lancamento ou se a extracao falhar - sem invalidar a captura
da Mensagem em si.

Cada execucao APPENDA linhas ao arquivo (ou cria com cabecalho em negrito
se nao existir). Idempotente apenas em serie - rodar um robo por vez.
"""
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from .relatorio import salvar_workbook_com_retry, sanitize_xlsx, carregar_workbook_seguro


_CABECALHO = [
    "CIV", "Cliente", "CNJ", "Estado", "Mensagem",
    "Data_Pagamento", "Saldo_Atualizado", "Capturado_em",
]


@dataclass
class LinhaMensagem:
    civ: str
    cliente: str
    cnj: str
    estado: str
    mensagem: str
    data_pagamento: str = ""
    saldo_atualizado: str = ""
    capturado_em: str = ""


def _garantir_workbook(caminho: Path):
    wb = carregar_workbook_seguro(caminho)
    if wb is not None:
        return wb, wb.active

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mensagens"
    ws.append(_CABECALHO)
    for col_idx in range(1, len(_CABECALHO) + 1):
        ws.cell(row=1, column=col_idx).font = Font(bold=True)
    return wb, ws


def registrar_mensagem(caminho_xlsx: Path, linha: LinhaMensagem) -> None:
    """Adiciona uma linha a planilha de saida (cria o arquivo se nao existir)."""
    caminho_xlsx = Path(caminho_xlsx)
    caminho_xlsx.parent.mkdir(parents=True, exist_ok=True)

    if not linha.capturado_em:
        linha.capturado_em = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    wb, ws = _garantir_workbook(caminho_xlsx)
    ws.append([
        sanitize_xlsx(linha.civ),
        sanitize_xlsx(linha.cliente),
        sanitize_xlsx(linha.cnj),
        sanitize_xlsx(linha.estado),
        sanitize_xlsx(linha.mensagem),
        sanitize_xlsx(linha.data_pagamento),
        sanitize_xlsx(linha.saldo_atualizado),
        sanitize_xlsx(linha.capturado_em),
    ])
    salvar_workbook_com_retry(wb, caminho_xlsx)
    wb.close()


def ja_capturado(caminho_xlsx: Path, civ: str) -> bool:
    """Idempotencia em UM arquivo: True se a CIV ja tem linha registrada
    no caminho dado. Util pra pular CIVs ja processadas em re-runs do
    MESMO arquivo. Para checar HISTORICO INTEIRO da pasta saida/ use
    ja_capturado_em_pasta()."""
    caminho_xlsx = Path(caminho_xlsx)
    if not caminho_xlsx.exists():
        return False
    civ_norm = (civ or "").strip().upper()
    if not civ_norm:
        return False
    wb = openpyxl.load_workbook(str(caminho_xlsx), read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            cel = row[0]
            if cel and str(cel).strip().upper() == civ_norm:
                return True
    finally:
        wb.close()
    return False


def ja_capturado_em_pasta(pasta_saida: Path, civ: str) -> bool:
    """Idempotencia CROSS-ARQUIVOS: True se a CIV foi capturada em
    QUALQUER planilha mensagens_acordo*.xlsx da pasta. Usado quando
    cada rodada gera arquivo datado novo - mantem idempotencia mesmo
    sem arquivo fixo de referencia.

    Procura por padrao 'mensagens_acordo*.xlsx' (incluindo o legado
    'mensagens_acordo.xlsx' sem timestamp).
    """
    pasta_saida = Path(pasta_saida)
    if not pasta_saida.exists():
        return False
    civ_norm = (civ or "").strip().upper()
    if not civ_norm:
        return False
    arquivos = sorted(pasta_saida.glob("mensagens_acordo*.xlsx"))
    for arq in arquivos:
        if arq.name.startswith("~$"):
            continue  # lock do Office
        try:
            if ja_capturado(arq, civ):
                return True
        except Exception:
            continue
    return False
