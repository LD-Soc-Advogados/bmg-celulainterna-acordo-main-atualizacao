"""
Leitor da planilha 'Planilha Envio de casos - Cedula Interna XXXX.xlsx'
mandada pelo Banco BMG (CIVs autorizadas a acordo).

Layout esperado (confirmado no arquivo de 19/05):
  Linha 1: headers (Cod. Causa | Workflow: Instalacao (Data) | Workflow:
           Status | Prioridade Causa - Descricao | Recurso | UF | Parte
           CON Principal - Nome | Numeracao unica | Resposavel | Observacoes)
  Linha 2+: dados

Diferente do `ler_planilha_qmc` do bmg-workflow-closer: nao exige coluna
'Workflow: Tarefa (Nome)' (essa planilha nao tem - todas as CIVs ja vem
filtradas como 'autorizadas a acordo' pelo Banco).
"""
from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata

import openpyxl


@dataclass
class CasoAutorizado:
    civ: str               # CIV1387277 (canonica com prefixo)
    civ_num: str           # 1387277 (so digitos)
    cliente: str
    cnj: str
    recurso: str           # escritorio adversario
    uf: str
    responsavel: str       # operador interno (Louis/Andreza/Elias)
    observacoes: str


_MAPA_COLUNAS = {
    "civ":          ["cod. causa", "cod causa", "codigo causa",
                     "codigo da causa", "civ"],
    "cliente":      ["parte con principal - nome",
                     "parte con principal-nome",
                     "cliente", "nome"],
    # "no processo" = "Nº Processo" depois da normalizacao NFKD (º -> o).
    # Header usado na planilha Consolidado_CedulaInterna_BMG_*.xlsx.
    "cnj":          ["numeracao unica", "numero unico", "cnj",
                     "no processo", "n processo", "n. processo",
                     "numero processo", "numero do processo", "processo"],
    "recurso":      ["recurso", "escritorio", "adversario"],
    "uf":           ["uf", "estado"],
    # 'resposavel' (sic) e o que vem na planilha real - aceito as duas grafias
    "responsavel":  ["resposavel", "responsavel"],
    "observacoes":  ["observacoes", "observacao", "obs"],
}

_RE_CIV_PREFIXO = re.compile(r"^CIV\s*\d{4,10}$", re.I)
_RE_CIV_NUMERO = re.compile(r"^\d{6,8}$")


def _normaliza(s) -> str:
    if s is None:
        return ""
    s = str(s)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.strip().lower()


def _normaliza_civ(valor) -> tuple:
    s = (str(valor) if valor is not None else "").strip().upper().replace(" ", "")
    if s.startswith("CIV"):
        numero = s[3:]
    else:
        numero = s
    return (f"CIV{numero}", numero)


def _mapear_headers(headers_linha) -> dict:
    if not headers_linha:
        return {}
    norm = [_normaliza(h) for h in headers_linha]
    posicoes = {}
    for chave, variantes in _MAPA_COLUNAS.items():
        for var in variantes:
            try:
                posicoes[chave] = norm.index(var)
                break
            except ValueError:
                continue
    return posicoes


def _eh_linha_vazia(linha) -> bool:
    return not linha or all(
        c is None or (isinstance(c, str) and not c.strip()) for c in linha
    )


def _eh_linha_civ_valida(linha, idx_civ) -> bool:
    if idx_civ is None or idx_civ >= len(linha):
        return False
    val = linha[idx_civ]
    if val is None:
        return False
    s = str(val).strip()
    return bool(_RE_CIV_PREFIXO.match(s) or _RE_CIV_NUMERO.match(s))


def ler_planilha_casos(caminho: Path) -> list:
    """Le a planilha de CIVs autorizadas e devolve lista de CasoAutorizado.

    Detecta o header automaticamente nas primeiras 5 linhas. Requer ao
    menos as colunas 'civ' e 'cliente' (as outras sao opcionais; se
    faltar alguma, vem como string vazia no dataclass).

    Levanta:
      FileNotFoundError se o arquivo nao existe.
      ValueError se nao conseguiu localizar 'civ' e 'cliente' nas
        primeiras 5 linhas.
    """
    caminho = Path(caminho)
    if not caminho.exists():
        raise FileNotFoundError(f"Planilha nao encontrada: {caminho}")

    # Planilha aberta no Excel (ou lock do OneDrive) da PermissionError
    # mesmo em leitura. Fallback: copia pra temp e le a copia - copiar
    # um arquivo aberto no Excel funciona, abrir direto nao.
    copia_temp = None
    try:
        wb = openpyxl.load_workbook(str(caminho), read_only=True, data_only=True)
    except PermissionError:
        import shutil
        import tempfile
        fd, tmp = tempfile.mkstemp(suffix=".xlsx")
        import os
        os.close(fd)
        copia_temp = Path(tmp)
        shutil.copy2(str(caminho), str(copia_temp))
        print(">>> Planilha bloqueada (aberta no Excel?) - lendo por copia temporaria.")
        wb = openpyxl.load_workbook(str(copia_temp), read_only=True, data_only=True)

    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
        if copia_temp is not None:
            try:
                copia_temp.unlink()
            except Exception:
                pass

    if not rows:
        return []

    posicoes = {}
    data_start = None
    obrigatorias = {"civ", "cliente"}

    for idx in range(min(5, len(rows))):
        if _eh_linha_vazia(rows[idx]):
            continue
        cand = _mapear_headers(rows[idx])
        if obrigatorias.issubset(cand.keys()):
            posicoes = cand
            data_start = idx + 1
            break

    if not posicoes:
        raise ValueError(
            "Nao consegui identificar as colunas da planilha. "
            "Espera-se header com 'Cod. Causa' (ou variante) e "
            "'Parte CON Principal - Nome' (ou 'Cliente') nas primeiras 5 linhas."
        )

    print(f">>> Planilha: header detectado na linha {data_start}, "
          f"colunas: {sorted(posicoes.keys())}")

    casos = []
    for row in rows[data_start:]:
        if _eh_linha_vazia(row):
            continue
        if not _eh_linha_civ_valida(row, posicoes.get("civ")):
            continue

        civ_canon, civ_num = _normaliza_civ(row[posicoes["civ"]])

        def _pega(chave):
            idx = posicoes.get(chave)
            if idx is None or idx >= len(row):
                return ""
            val = row[idx]
            return str(val).strip() if val is not None else ""

        casos.append(CasoAutorizado(
            civ=civ_canon,
            civ_num=civ_num,
            cliente=_pega("cliente"),
            cnj=_pega("cnj"),
            recurso=_pega("recurso"),
            uf=_pega("uf"),
            responsavel=_pega("responsavel"),
            observacoes=_pega("observacoes"),
        ))

    return casos
